"""Turn an uploaded resume (PDF or Word) into resume JSON.

Usage: py extract.py <resume.pdf or resume.docx> <output.json>

Steps:
  1. Read the text out of the file.
  2. AI call 1: extract the main resume. Sections the schema doesn't cover
     come back in "unplaced", grouped by heading.
  3. AI call 2 (once per unplaced group): turn it into a custom section.
  4. Add the custom sections to the resume, in their original position.
  5. Remove empty values, check against the full schema, and save.
"""
import json
import re
import sys
from pathlib import Path

import anthropic
import docx
import jsonschema
import pdfplumber
from dotenv import load_dotenv

import columns

load_dotenv()
client = anthropic.Anthropic()

MODEL = "claude-haiku-5-5"
HERE = Path(__file__).parent
EXTRACTION_INSTRUCTIONS = (HERE / "prompts" / "extraction.md").read_text(encoding="utf-8")
CUSTOM_SECTION_INSTRUCTIONS = (HERE / "prompts" / "custom_section.md").read_text(encoding="utf-8")

with open(HERE / ".." / "resume.schema.json", encoding="utf-8") as f:
    RESUME_SCHEMA = json.load(f)

BUILTIN_SECTIONS = ["summary", "education", "experience", "projects", "skills"]


NOT_SUPPORTED = {"maxLength", "minLength", "maxItems", "uniqueItems", "pattern",
                 "$schema", "$id", "format"}


def simplify(node):
    """Copy a schema for the AI: keep $refs (each definition is written once),
    remove rules the API can't use, and make every field required
    (the AI writes "" or [] when a field doesn't apply)."""
    if isinstance(node, list):
        return [simplify(item) for item in node]
    if not isinstance(node, dict):
        return node
    if "$ref" in node:
        return {"$ref": node["$ref"]}
    if "anyOf" in node:
        return {"type": "string"}
    result = {}
    for key, value in node.items():
        if key in NOT_SUPPORTED or key == "required":
            continue
        if key == "minItems" and value > 1:
            continue
        if key in ("properties", "$defs"):
            result[key] = {name: simplify(sub) for name, sub in value.items()}
            if key == "properties":
                result["required"] = list(value.keys())
        else:
            result[key] = simplify(value)
    return result


def build_resume_wrapper():
    """Schema for AI call 1: the resume (without custom sections, which made
    the schema too large) plus the unplaced groups."""
    resume_part = simplify(RESUME_SCHEMA)
    del resume_part["properties"]["custom_sections"]
    resume_part["required"].remove("custom_sections")

    defs = resume_part.pop("$defs")
    for name in ("custom_section", "custom_entry", "custom_id"):
        del defs[name]

    unplaced_group = {
        "type": "object",
        "properties": {
            "heading": {"type": "string"},
            "lines": {"type": "array", "items": {"type": "string"}},
            "after": {"type": "string"},
        },
        "required": ["heading", "lines", "after"],
        "additionalProperties": False,
    }

    return {
        "type": "object",
        "$defs": defs,
        "properties": {
            "resume": resume_part,
            "unplaced": {"type": "array", "items": unplaced_group},
        },
        "required": ["resume", "unplaced"],
        "additionalProperties": False,
    }


def build_custom_section_schema():
    """Schema for AI call 2: one custom section's layout and content.
    No id or title: the code sets those from the heading."""
    all_defs = simplify(RESUME_SCHEMA)["$defs"]
    section = all_defs["custom_section"]
    for field in ("id", "title"):
        del section["properties"][field]
        section["required"].remove(field)

    # Only the definitions this schema points to.
    needed = ("custom_entry", "bullet", "text", "date", "skill_group")
    section["$defs"] = {name: all_defs[name] for name in needed}
    return section


RESUME_WRAPPER = build_resume_wrapper()
CUSTOM_SECTION_SCHEMA = build_custom_section_schema()



def ask_ai(instructions, content, schema):
    """Send one request and return the reply as a dictionary."""
    response = client.messages.create(
        model=MODEL,
        max_tokens=15000,
        system=instructions,
        messages=[{"role": "user", "content": content}],
        output_config={"format": {"type": "json_schema", "schema": schema}},
    )

    if response.stop_reason == "max_tokens":
        raise RuntimeError("the reply was cut off; raise max_tokens")
    if response.stop_reason == "refusal":
        raise RuntimeError("the model declined to process this document")

    for block in response.content:
        if block.type == "text":
            return json.loads(block.text)

    raise RuntimeError("the reply had no text")


def extract_resume(text):
    """AI call 1: the main resume, plus unplaced groups."""
    return ask_ai(EXTRACTION_INSTRUCTIONS, f"<resume>\n{text}\n</resume>", RESUME_WRAPPER)


def extract_custom_section(group):
    """AI call 2: one unplaced group as a custom section's layout and content."""
    lines = "\n".join(group["lines"])
    content = f"<section>\nHeading: {group['heading']}\n{lines}\n</section>"
    return ask_ai(CUSTOM_SECTION_INSTRUCTIONS, content, CUSTOM_SECTION_SCHEMA)


def make_custom_id(heading, used_ids):
    """'Volunteer Experience' -> 'custom-volunteer-experience', unique."""
    slug = re.sub(r"[^a-z0-9]+", "-", heading.lower()).strip("-") or "section"
    candidate = f"custom-{slug}"
    number = 2
    while candidate in used_ids:
        candidate = f"custom-{slug}-{number}"
        number += 1
    return candidate


def insert_after(order, item, after):
    """Put item into the order list just after `after`.
    "" means first; an unknown `after` means last."""
    if after == "":
        order.insert(0, item)
    elif after in order:
        order.insert(order.index(after) + 1, item)
    else:
        order.append(item)


def add_custom_sections(resume, unplaced):
    """AI call 2 for each unplaced group, then add the results to the resume.
    Returns the titles of the sections added."""
    custom_sections = resume.setdefault("custom_sections", [])
    order = resume.setdefault("section_order", [])
    used_ids = {section["id"] for section in custom_sections}
    ids_by_heading = {}
    added = []

    for group in unplaced:
        heading = group["heading"].strip()
        if heading.isupper():
            heading = heading.title()
        content = extract_custom_section(group)

        section_id = make_custom_id(heading, used_ids)
        used_ids.add(section_id)
        ids_by_heading[heading.lower()] = section_id

        custom_sections.append({"id": section_id, "title": heading, **content})

        after = group["after"].strip()
        after = ids_by_heading.get(after.lower(), after.lower())
        insert_after(order, section_id, after)
        added.append(heading)

    return added



def remove_empty(value):
    if isinstance(value, dict):
        cleaned = {k: remove_empty(v) for k, v in value.items()}
        return {k: v for k, v in cleaned.items() if v not in ("", [], {})}
    if isinstance(value, list):
        cleaned = [remove_empty(v) for v in value]
        return [v for v in cleaned if v not in ("", [], {})]
    if isinstance(value, str):
        return value.strip()
    return value


def validate_resume(resume):
    validator = jsonschema.Draft202012Validator(
        RESUME_SCHEMA, format_checker=jsonschema.FormatChecker())
    problems = []
    for error in validator.iter_errors(resume):
        where = "/".join(str(part) for part in error.absolute_path) or "(top level)"
        message = error.message if len(error.message) <= 120 else error.message[:117] + "..."
        problems.append(f"{where}: {message}")

    custom_ids = [section.get("id") for section in resume.get("custom_sections", [])]
    if len(custom_ids) != len(set(custom_ids)):
        problems.append("custom_sections: two sections have the same id")
    for name in resume.get("section_order", []):
        if name.startswith("custom-") and name not in custom_ids:
            problems.append(f"section_order: '{name}' has no matching custom section")
        elif not name.startswith("custom-") and name not in BUILTIN_SECTIONS:
            problems.append(f"section_order: unknown section '{name}'")
    for section_id in custom_ids:
        if section_id not in resume.get("section_order", []):
            problems.append(f"custom_sections: '{section_id}' is missing from section_order")
    return problems


def extract_pdf_text(path):
    pages_text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages_text.append(columns.page_text(page))
    return "\n".join(pages_text)


def extract_docx_text(path):
    parts = []
    doc = docx.Document(path)
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join([cell.text for cell in row.cells]))
    return "\n".join(parts)


def fail(message):
    print(f"error: {message}", file=sys.stderr)
    sys.exit(1)


def main():
    if len(sys.argv) != 3:
        fail("usage: py extract.py <resume.pdf or resume.docx> <output.json>")

    input_path, output_path = sys.argv[1], sys.argv[2]

    if not Path(input_path).exists():
        fail(f"file not found: {input_path}")

    suffix = Path(input_path).suffix.lower()
    if suffix not in (".pdf", ".docx"):
        fail(f"unsupported file type '{suffix}': please upload a PDF or Word (.docx) file")

    try:
        if suffix == ".pdf":
            text = extract_pdf_text(input_path)
        else:
            text = extract_docx_text(input_path)
    except Exception as e:
        fail(f"could not read {input_path} ({e})")

    if len(text.strip()) < 100:
        fail("no readable text found. This may be a scanned image; "
             "please upload a text-based PDF or a Word file.")

    try:
        data = extract_resume(text)
        resume = remove_empty(data["resume"])
        unplaced = [group for group in data["unplaced"] if group["lines"]]
        added = add_custom_sections(resume, unplaced)
        resume = remove_empty(resume)
    except Exception as e:
        fail(f"extraction failed ({e})")

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(resume, f, indent=2, ensure_ascii=False)
    print(f"Wrote: {output_path}")

    if added:
        print("Added as custom sections: " + ", ".join(added))

    problems = validate_resume(resume)
    if problems:
        print("Warning: the result doesn't fully match the schema:", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)


if __name__ == "__main__":
    main()